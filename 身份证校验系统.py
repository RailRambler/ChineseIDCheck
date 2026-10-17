import tkinter as tk
from tkinter import ttk, messagebox
import re
from datetime import datetime
import os
import sys
from openpyxl import load_workbook

class IDCardCheckerApp:
    def __init__(self, root):
        self.root = root
        self.root.title("身份证智能分析系统")
        screen_width = self.root.winfo_screenwidth()
        screen_height = self.root.winfo_screenheight()
        win_width = int(screen_width * 0.5)
        win_height = int(screen_height * 0.6)
        if win_width < 700: win_width = 700
        if win_height < 680: win_height = 680
        self.root.geometry(f"{win_width}x{win_height}")
        self.root.minsize(700, 680)
        self.root.resizable(True, True)
        self.dark_mode = False
        self.setup_colors()
        self.setup_styles()
        self.weight_factors = [7, 9, 10, 5, 8, 4, 2, 1, 6, 3, 7, 9, 10, 5, 8, 4, 2]
        self.check_codes = ['1', '0', 'X', '9', '8', '7', '6', '5', '4', '3', '2']
        self.address_code_map = {}
        self.province_code_map = {}
        self.load_address_data()
        self.placeholder_text = "请输入身份证号码"
        self.has_content = False
        self.current_result_data = None
        self.show_id = False
        self.setup_ui()

    def resource_path(self, relative_path):
        if hasattr(sys, '_MEIPASS'):
            return os.path.join(sys._MEIPASS, relative_path)
        return os.path.join(os.path.abspath("."), relative_path)

    def setup_colors(self):
        if self.dark_mode:
            self.colors = {
                'primary': '#0A84FF',
                'primary_dark': '#007AFF',
                'secondary': '#636366',
                'success': '#30D158',
                'danger': '#FF453A',
                'warning': '#FFD60A',
                'bg': '#000000',
                'card_bg': '#2c2c2e',
                'text_light': '#f5f5f7',
                'text_muted': '#86868b',
                'cursor': '#f5f5f7',
                'entry_fg': '#f5f5f7',
                'entry_placeholder': '#636366',
                'border': '#3a3a3c',
                'button_bg': '#0A84FF',
                'button_fg': 'white',
                'button_hover': '#007AFF',
                'secondary_bg': '#636366',
                'secondary_fg': 'white',
                'secondary_hover': '#4a4a4d',
                'theme_btn_fg': '#aaaaaa',
                'theme_btn_hover': '#f5f5f7',
                'icon_color': '#f5f5f7'
            }
        else:
            self.colors = {
                'primary': '#007AFF',
                'primary_dark': '#0055CC',
                'secondary': '#8e8e93',
                'success': '#34C759',
                'danger': '#FF3B30',
                'warning': '#FF9500',
                'bg': '#f5f5f7',
                'card_bg': '#ffffff',
                'text_light': '#1d1d1f',
                'text_muted': '#86868b',
                'cursor': '#1d1d1f',
                'entry_fg': '#1d1d1f',
                'entry_placeholder': '#8e8e93',
                'border': '#d2d2d7',
                'button_bg': '#007AFF',
                'button_fg': 'white',
                'button_hover': '#0055CC',
                'secondary_bg': '#8e8e93',
                'secondary_fg': 'white',
                'secondary_hover': '#6e6e73',
                'theme_btn_fg': '#86868b',
                'theme_btn_hover': '#1d1d1f',
                'icon_color': '#1d1d1f'
            }

    def setup_styles(self):
        style = ttk.Style()
        style.theme_use('clam')
        font_family = "-apple-system, BlinkMacSystemFont, 'Segoe UI', Roboto, Helvetica, Arial, sans-serif"
        style.configure('Modern.TFrame', background=self.colors['bg'])
        style.configure('Card.TFrame', background=self.colors['card_bg'], relief='flat')
        style.configure('Title.TLabel',
                       font=(font_family, 28, 'bold'),
                       foreground=self.colors['text_light'],
                       background=self.colors['bg'])
        style.configure('Subtitle.TLabel',
                       font=(font_family, 13),
                       foreground=self.colors['text_muted'],
                       background=self.colors['bg'])
        style.configure('CardTitle.TLabel',
                       font=(font_family, 16, 'bold'),
                       foreground=self.colors['text_light'],
                       background=self.colors['bg'])
        style.configure('Value.TLabel',
                       font=(font_family, 15),
                       foreground=self.colors['text_light'],
                       background=self.colors['bg'])
        style.configure('Info.TLabel',
                       font=(font_family, 15),
                       foreground=self.colors['text_muted'],
                       background=self.colors['bg'])
        style.configure('Copy.TButton',
                       font=(font_family, 13),
                       background=self.colors['primary'],
                       foreground='white',
                       borderwidth=0,
                       focusthickness=0,
                       padding=(15, 8))
        style.map('Copy.TButton',
                 background=[('active', self.colors['primary_dark'])])
        style.configure('Action.TButton',
                       font=(font_family, 18, 'bold'),
                       background=self.colors['button_bg'],
                       foreground=self.colors['button_fg'],
                       borderwidth=2,
                       relief='raised',
                       padding=(20, 10),
                       focuscolor='',
                       highlightthickness=0)
        style.map('Action.TButton',
                 background=[('active', self.colors['button_hover'])])
        style.configure('Secondary.TButton',
                       font=(font_family, 18),
                       background=self.colors['secondary_bg'],
                       foreground=self.colors['secondary_fg'],
                       borderwidth=2,
                       relief='raised',
                       padding=(20, 10),
                       focuscolor='',
                       highlightthickness=0)
        style.map('Secondary.TButton',
                 background=[('active', self.colors['secondary_hover'])])

    def setup_ui(self):
        self.root.configure(bg=self.colors['bg'])
        main_container = ttk.Frame(self.root, style='Modern.TFrame')
        main_container.pack(fill='both', expand=True, padx=40, pady=30)
        self.create_header(main_container)
        self.create_input_section(main_container)
        self.create_result_section(main_container)

    def create_header(self, parent):
        header_frame = ttk.Frame(parent, style='Modern.TFrame')
        header_frame.pack(fill='x', pady=(0, 20))
        top_frame = ttk.Frame(header_frame, style='Modern.TFrame')
        top_frame.pack(fill='x', pady=(0, 12))
        theme_text = "🌙 深色" if self.dark_mode else "☀️ 浅色"
        self.theme_button = tk.Button(
            top_frame,
            text=theme_text,
            command=self.toggle_theme,
            font=('-apple-system', 13),
            relief='flat',
            bd=0,
            anchor='center',
            justify='center',
            highlightthickness=0,
            bg=self.colors['bg'],
            fg=self.colors['theme_btn_fg'],
            activebackground=self.colors['bg'],
            activeforeground=self.colors['theme_btn_hover'],
            cursor='hand2'
        )
        self.theme_button.pack(side='right', padx=(0, 10), pady=5)
        title_container = ttk.Frame(header_frame, style='Modern.TFrame')
        title_container.pack(anchor='center')
        self.icon_label = tk.Label(title_container,
                                   text="🪪",
                                   font=('Segoe UI Emoji', 36),
                                   bg=self.colors['bg'],
                                   fg=self.colors['primary'])
        self.icon_label.pack()
        title_label = ttk.Label(title_container,
                               text="身份证智能分析",
                               style='Title.TLabel')
        title_label.pack(pady=(8, 2))
        subtitle_label = ttk.Label(title_container,
                                  text="智能验证 · 户籍分析 · 信息提取",
                                  style='Subtitle.TLabel')
        subtitle_label.pack()

    def create_input_section(self, parent):
        input_frame = ttk.Frame(parent, style='Modern.TFrame')
        input_frame.pack(fill='x', pady=(0, 18))
        ttk.Label(input_frame,
                 text="请输入身份证号码",
                 style='CardTitle.TLabel').pack(anchor='center', pady=(0, 10))
        input_group = ttk.Frame(input_frame, style='Modern.TFrame')
        input_group.pack(fill='x', padx=10)

        self.id_entry = tk.Entry(
            input_group,
            font=('Helvetica Neue', 18),
            bg=self.colors['card_bg'],
            fg=self.colors['entry_placeholder'],
            insertbackground=self.colors['cursor'],
            relief='solid',
            bd=2,
            highlightthickness=2,
            highlightcolor=self.colors['primary'],
            highlightbackground=self.colors['border'],
            show=""
        )
        self.id_entry.pack(side='left', fill='x', expand=True, padx=(0, 5), ipady=8)
        self.id_entry.insert(0, self.placeholder_text)

        def on_focus_in(e):
            current = self.id_entry.get()
            if current == self.placeholder_text:
                self.id_entry.delete(0, tk.END)
                self.id_entry.config(fg=self.colors['entry_fg'])
                if self.show_id:
                    self.id_entry.config(show="")
                else:
                    self.id_entry.config(show="*")
            self.id_entry.config(highlightbackground=self.colors['primary'])
            self.update_show_button_visibility()

        def on_focus_out(e):
            current = self.id_entry.get().strip()
            if not current:
                self.id_entry.delete(0, tk.END)
                self.id_entry.insert(0, self.placeholder_text)
                self.id_entry.config(fg=self.colors['entry_placeholder'])
                self.id_entry.config(show="")
                self.show_id = False
                self.toggle_show_btn.config(text="👁️")
            self.id_entry.config(highlightbackground=self.colors['border'])
            self.update_show_button_visibility()

        def on_key_release(e):
            self.update_show_button_visibility()

        self.id_entry.bind('<FocusIn>', on_focus_in)
        self.id_entry.bind('<FocusOut>', on_focus_out)
        self.id_entry.bind('<KeyRelease>', on_key_release)
        self.id_entry.bind('<Return>', lambda e: self.process_id_card())

        btn_group = ttk.Frame(input_group, style='Modern.TFrame')
        btn_group.pack(side='left', padx=(0, 0))

        self.toggle_show_btn = tk.Button(
            btn_group,
            text="👁️",
            font=('Segoe UI Emoji', 16),
            relief='flat',
            bd=0,
            highlightthickness=0,
            bg=self.colors['bg'],
            fg=self.colors['icon_color'],
            activebackground=self.colors['bg'],
            activeforeground=self.colors['icon_color'],
            cursor='hand2',
            command=self.toggle_show
        )

        self.verify_btn = ttk.Button(
            btn_group,
            text="分析",
            style='Action.TButton',
            command=self.process_id_card
        )
        self.clear_btn = ttk.Button(
            btn_group,
            text="清空",
            style='Secondary.TButton',
            command=self.clear_all
        )

        self.verify_btn.pack(side='left', padx=5)
        self.clear_btn.pack(side='left', padx=5)

        self.update_show_button_visibility()

    def update_show_button_visibility(self):
        current = self.id_entry.get()
        if current and current != self.placeholder_text:
            try:
                if not self.toggle_show_btn.winfo_ismapped():
                    self.verify_btn.pack_forget()
                    self.clear_btn.pack_forget()
                    self.toggle_show_btn.pack(side='left', padx=5)
                    self.verify_btn.pack(side='left', padx=5)
                    self.clear_btn.pack(side='left', padx=5)
            except:
                pass
        else:
            try:
                if self.toggle_show_btn.winfo_ismapped():
                    self.toggle_show_btn.pack_forget()
            except:
                pass

    def toggle_show(self):
        current = self.id_entry.get()
        if current == self.placeholder_text or not current:
            return
        self.show_id = not self.show_id
        if self.show_id:
            self.id_entry.config(show="")
            self.toggle_show_btn.config(text="🙈")
        else:
            self.id_entry.config(show="*")
            self.toggle_show_btn.config(text="👁️")

    def create_result_section(self, parent):
        container = ttk.Frame(parent, style='Modern.TFrame')
        container.pack(fill='both', expand=True)
        self.result_canvas = tk.Canvas(container, bg=self.colors['bg'], highlightthickness=0)
        scrollbar = ttk.Scrollbar(container, orient='vertical', command=self.result_canvas.yview)
        self.result_canvas.configure(yscrollcommand=scrollbar.set)
        self.scrollable_frame = ttk.Frame(self.result_canvas, style='Modern.TFrame')
        self.scrollable_frame.bind(
            '<Configure>',
            lambda e: self.result_canvas.configure(scrollregion=self.result_canvas.bbox('all'))
        )
        self.canvas_window = self.result_canvas.create_window((0, 0), window=self.scrollable_frame, anchor='nw', width=self.result_canvas.winfo_width())

        def on_canvas_configure(e):
            self.result_canvas.itemconfig(self.canvas_window, width=e.width)
        self.result_canvas.bind('<Configure>', on_canvas_configure)
        self.result_canvas.pack(side='left', fill='both', expand=True)
        scrollbar.pack(side='right', fill='y')
        self.show_initial_state()

    def toggle_theme(self):
        self.dark_mode = not self.dark_mode
        self.setup_colors()
        self.setup_styles()
        self.refresh_ui()

    def refresh_ui(self):
        self.root.configure(bg=self.colors['bg'])
        theme_text = "🌙 深色" if self.dark_mode else "☀️ 浅色"
        self.theme_button.config(
            text=theme_text,
            bg=self.colors['bg'],
            fg=self.colors['theme_btn_fg'],
            activebackground=self.colors['bg'],
            activeforeground=self.colors['theme_btn_hover']
        )
        self.icon_label.config(fg=self.colors['primary'], bg=self.colors['bg'])
        entry = self.id_entry
        current_text = entry.get()
        if current_text == self.placeholder_text:
            entry.config(bg=self.colors['card_bg'], fg=self.colors['entry_placeholder'])
            entry.config(show="")
        else:
            entry.config(bg=self.colors['card_bg'], fg=self.colors['entry_fg'])
            if self.show_id:
                entry.config(show="")
            else:
                entry.config(show="*")
        entry.config(insertbackground=self.colors['cursor'],
                     highlightcolor=self.colors['primary'],
                     highlightbackground=self.colors['border'])
        self.setup_styles()
        self.result_canvas.config(bg=self.colors['bg'])
        self.scrollable_frame.configure(style='Modern.TFrame')
        self.toggle_show_btn.config(bg=self.colors['bg'], fg=self.colors['icon_color'],
                                    activebackground=self.colors['bg'],
                                    activeforeground=self.colors['icon_color'])
        self.update_show_button_visibility()
        if self.current_result_data:
            address, birth_date, gender, check_result, full_id = self.current_result_data
            self.show_results(address, birth_date, gender, check_result, full_id)
        else:
            self.show_initial_state()

    def show_initial_state(self):
        for widget in self.scrollable_frame.winfo_children():
            widget.destroy()
        wrapper = ttk.Frame(self.scrollable_frame, style='Modern.TFrame')
        wrapper.pack(fill='both', expand=True)
        icon_label = tk.Label(wrapper,
                             text="🪪",
                             font=('Segoe UI Emoji', 64),
                             bg=self.colors['bg'],
                             fg=self.colors['text_muted'])
        icon_label.pack(pady=(30, 15))
        title_label = ttk.Label(wrapper,
                               text="等待输入身份证号码",
                               style='Title.TLabel')
        title_label.pack(pady=(0, 8))
        hint_label = ttk.Label(wrapper,
                              text="输入17位生成校验码，输入18位验证正确性",
                              style='Subtitle.TLabel')
        hint_label.pack()

    def add_info_row(self, parent, label, value, icon="", color=None):
        row = ttk.Frame(parent, style='Modern.TFrame')
        row.pack(fill='x', pady=4, padx=20)
        left_frame = ttk.Frame(row, style='Modern.TFrame')
        left_frame.pack(side='left')
        if icon:
            icon_lbl = tk.Label(left_frame,
                               text=icon,
                               font=('Segoe UI Emoji', 16),
                               bg=self.colors['bg'],
                               fg=color if color else self.colors['primary'])
            icon_lbl.pack(side='left', padx=(0, 8))
        lbl = ttk.Label(left_frame, text=label, style='Info.TLabel')
        lbl.pack(side='left')
        val = ttk.Label(row, text=value, style='Value.TLabel')
        val.pack(side='right')
        return row

    def show_results(self, address, birth_date, gender, check_result, full_id=None):
        for widget in self.scrollable_frame.winfo_children():
            widget.destroy()
        self.current_result_data = (address, birth_date, gender, check_result, full_id)
        wrapper = ttk.Frame(self.scrollable_frame, style='Modern.TFrame')
        wrapper.pack(fill='both', expand=True)
        ttk.Label(wrapper, text="基本信息", style='CardTitle.TLabel',
                  foreground=self.colors['primary']).pack(anchor='center', pady=(5, 10))
        self.add_info_row(wrapper, "户籍地区", address, "📍", self.colors['success'])
        self.add_info_row(wrapper, "出生日期", birth_date, "📅", self.colors['primary'])
        self.add_info_row(wrapper, "性别", gender, "👤", self.colors['warning'])
        ttk.Separator(wrapper, orient='horizontal').pack(fill='x', pady=10, padx=40)
        ttk.Label(wrapper, text="验证结果", style='CardTitle.TLabel',
                  foreground=self.colors['primary']).pack(anchor='center', pady=(5, 10))
        is_error = "失败" in check_result
        result_icon = "❌" if is_error else "✅"
        color = self.colors['danger'] if is_error else self.colors['success']
        self.add_info_row(wrapper, "", check_result, result_icon, color)
        if full_id:
            ttk.Separator(wrapper, orient='horizontal').pack(fill='x', pady=10, padx=40)
            ttk.Label(wrapper, text="生成结果", style='CardTitle.TLabel',
                      foreground=self.colors['primary']).pack(anchor='center', pady=(5, 10))
            formatted_id = full_id
            row = ttk.Frame(wrapper, style='Modern.TFrame')
            row.pack(fill='x', pady=4, padx=20)
            left_frame = ttk.Frame(row, style='Modern.TFrame')
            left_frame.pack(side='left')
            icon_lbl = tk.Label(left_frame,
                               text="🆔",
                               font=('Segoe UI Emoji', 16),
                               bg=self.colors['bg'],
                               fg=self.colors['primary'])
            icon_lbl.pack(side='left', padx=(0, 8))
            ttk.Label(left_frame, text="完整号码", style='Info.TLabel').pack(side='left')
            right_frame = ttk.Frame(row, style='Modern.TFrame')
            right_frame.pack(side='right')
            ttk.Label(right_frame, text=formatted_id, style='Value.TLabel').pack(side='left', padx=(0, 12))
            copy_btn = ttk.Button(right_frame,
                                 text="复制",
                                 style='Copy.TButton',
                                 command=lambda: self.copy_to_clipboard(full_id))
            copy_btn.pack(side='right')
        ttk.Separator(wrapper, orient='horizontal').pack(fill='x', pady=10, padx=40)
        time_label = ttk.Label(wrapper,
                              text=f"分析时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
                              style='Info.TLabel')
        time_label.pack(anchor='center', pady=(5, 0))

    def copy_to_clipboard(self, text):
        self.root.clipboard_clear()
        self.root.clipboard_append(text)
        self.root.update()
        messagebox.showinfo("复制成功", "身份证号码已复制到剪贴板！")

    def load_address_data(self):
        xlsx_path = self.resource_path("address_codes.xlsx")
        if not os.path.exists(xlsx_path):
            messagebox.showerror("错误", f"未找到地址码文件：{xlsx_path}\n请将 'address_codes.xlsx' 放在程序同目录下。")
            return
        try:
            wb = load_workbook(xlsx_path, data_only=True)
            sheet = wb.active
            for row in sheet.iter_rows(min_row=1, max_col=2, values_only=True):
                if row and len(row) >= 2:
                    name = str(row[0]).strip() if row[0] is not None else ""
                    code = str(row[1]).strip() if row[1] is not None else ""
                    if name and code:
                        self._process_row([name, code])
        except Exception as e:
            messagebox.showerror("错误", f"读取 Excel 文件失败：{str(e)}")

    def _process_row(self, row):
        if row is None or len(row) < 2:
            return
        name = row[0].strip()
        code = row[1].strip()
        if not code or not name:
            return
        if len(code) == 6:
            self.address_code_map[code] = name

    def get_address_info(self, id_card):
        if len(id_card) < 6:
            return "地址码不完整"
        address_code = id_card[:6]
        if address_code in self.address_code_map:
            return self.address_code_map[address_code]
        else:
            return "未知地区"

    def get_birth_info(self, id_card):
        if len(id_card) < 17:
            return "信息不完整", "未知"
        birth_year = id_card[6:10]
        birth_month = id_card[10:12]
        birth_day = id_card[12:14]
        try:
            birth_date = datetime.strptime(f"{birth_year}-{birth_month}-{birth_day}", "%Y-%m-%d")
            today = datetime.today()
            age = today.year - birth_date.year - ((today.month, today.day) < (birth_date.month, birth_date.day))
            birth_text = f"{birth_year}年{birth_month}月{birth_day}日 (年龄: {age}岁)"
        except:
            birth_text = f"{birth_year}年{birth_month}月{birth_day}日"
        gender_digit = int(id_card[16])
        gender = "男性" if gender_digit % 2 == 1 else "女性"
        return birth_text, gender

    def process_id_card(self):
        user_input = self.id_entry.get().strip().upper()
        if not user_input or user_input == self.placeholder_text:
            messagebox.showwarning("输入错误", "请输入身份证号！")
            return
        user_input = re.sub(r'[-\s]', '', user_input)
        if len(user_input) == 18:
            if not user_input[:17].isdigit():
                messagebox.showerror("输入错误", "前17位必须为数字！")
                return
            total = sum(int(user_input[i]) * self.weight_factors[i] for i in range(17))
            remainder = total % 11
            correct_check_code = self.check_codes[remainder]
            address = self.get_address_info(user_input)
            birth_date, gender = self.get_birth_info(user_input)
            if user_input[-1] == correct_check_code:
                check_result = f"✅ 验证通过 - 校验码正确 ({correct_check_code})"
            else:
                check_result = f"❌ 验证失败 - 应为 {correct_check_code}，输入 {user_input[-1]}"
            self.show_results(address, birth_date, gender, check_result)
        elif len(user_input) == 17:
            if not user_input.isdigit():
                messagebox.showerror("输入错误", "前17位必须为数字！")
                return
            total = sum(int(user_input[i]) * self.weight_factors[i] for i in range(17))
            remainder = total % 11
            generated_check_code = self.check_codes[remainder]
            address = self.get_address_info(user_input)
            birth_date, gender = self.get_birth_info(user_input)
            full_id = user_input + generated_check_code
            check_result = f"✅ 已生成校验码: {generated_check_code}"
            self.show_results(address, birth_date, gender, check_result, full_id)
        else:
            messagebox.showerror("输入错误", "输入长度必须为17位（生成校验码）或18位（验证正确性）！")

    def clear_all(self):
        self.id_entry.delete(0, tk.END)
        self.id_entry.insert(0, self.placeholder_text)
        self.id_entry.config(fg=self.colors['entry_placeholder'])
        self.id_entry.config(show="")
        self.show_id = False
        self.toggle_show_btn.config(text="👁️")
        self.has_content = False
        self.current_result_data = None
        self.update_show_button_visibility()
        self.show_initial_state()

def main():
    root = tk.Tk()
    app = IDCardCheckerApp(root)
    root.mainloop()

if __name__ == "__main__":
    main()
